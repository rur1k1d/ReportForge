"""Exporters: render a report to PDF (WeasyPrint), DOCX (python-docx),
and save the raw HTML for browser review.

PDF output is generated from the same Jinja2 HTML template used for preview,
so the two never drift apart.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from . import renderer
from .branding import SEVERITY_COLORS


def _rgb(hex_color: str) -> RGBColor:
    """'#B22222' -> RGBColor. Falls back to grey for bad input."""
    hex_color = hex_color.lstrip("#")
    if len(hex_color) != 6:
        return RGBColor(0x66, 0x66, 0x66)
    try:
        return RGBColor(int(hex_color[0:2], 16),
                        int(hex_color[2:4], 16),
                        int(hex_color[4:6], 16))
    except ValueError:
        return RGBColor(0x66, 0x66, 0x66)


def _write(path: Path, data: bytes | str, binary: bool = True) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    mode = "wb" if binary else "w"
    kwargs = {} if binary else {"encoding": "utf-8"}
    if isinstance(data, str):
        mode = "w"
        kwargs = {"encoding": "utf-8"}
    with open(path, mode, **kwargs) as fh:
        fh.write(data)
    return path


def render_html(report, findings, branding_ref: str = "default") -> str:
    return renderer.render_html(report, findings, branding_ref)


def export_pdf(report, findings, out: Path, branding_ref: str = "default") -> Path:
    """Render HTML then convert via WeasyPrint to a PDF file."""
    from weasyprint import HTML

    html = render_html(report, findings, branding_ref)
    out.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html).write_pdf(str(out))
    return out


def export_docx(report, findings, out: Path, branding_ref: str = "default") -> Path:
    """Minimal, dependency-light DOCX via python-docx.

    Not a pixel-perfect copy of the HTML: it produces the editable document
    structure clients actually ask for (title, exec summary, findings list).
    """
    doc = Document()
    title = doc.add_heading(report.title or "Pentest Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Client: {report.client}")
    doc.add_paragraph(f"Scope: {report.scope or '—'}")

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(report.summary or renderer.exec_summary_text(
        report, renderer.stats(findings)))

    doc.add_heading("Findings", level=1)
    for f in findings:
        p = doc.add_heading(f.title, level=2)
        color = _rgb(SEVERITY_COLORS.get(f.severity, "#666666"))
        p.runs[0].font.color.rgb = color
        meta = doc.add_paragraph()
        meta.add_run(f"Severity: {f.severity.upper()}")
        if f.cwe:
            meta.add_run(f"  CWE: {f.cwe}")
        if f.asset:
            meta.add_run(f"  Asset: {f.asset}")
        if f.description:
            doc.add_paragraph(f"Description: {f.description}")
        if f.impact:
            doc.add_paragraph(f"Impact: {f.impact}")
        if f.recommendation:
            doc.add_paragraph(f"Recommendation: {f.recommendation}")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def export_html(report, findings, out: Path, branding_ref: str = "default") -> Path:
    """Write the rendered HTML to disk (static view / pandoc input)."""
    html = render_html(report, findings, branding_ref)
    return _write(out, html)


# ---------------------------------------------------------------- markdown ---


def render_markdown(report, findings, branding_ref: str = "default") -> str:
    """Native Markdown rendering (no pandoc dependency).

    Mirrors the HTML report structure: title + meta, executive summary,
    severity counts, per-finding sections grouped high->info.
    """
    from .renderer import stats, exec_summary_text

    st = stats(findings)
    severity_order = ["critical", "high", "medium", "low", "info"]
    lines: list[str] = []

    lines.append(f"# {report.title}")
    lines.append("")
    lines.append(f"- **Client:** {report.client}")
    lines.append(f"- **Scope:** {report.scope or '—'}")
    lines.append(
        f"- **Template:** {report.template_type} | **Version:** {report.version}"
    )
    if getattr(report, "updated_at", None):
        lines.append(
            f"- **Generated:** {report.updated_at.strftime('%Y-%m-%d')}"
        )
    lines.append("")

    summary = (getattr(report, "summary", "") or "").strip() or exec_summary_text(
        report, st
    )
    lines.append("## Executive Summary")
    lines.append("")
    lines.append(summary)
    lines.append("")

    lines.append("## Findings Overview")
    lines.append("")
    lines.append("| Severity | Count |")
    lines.append("|---|---|")
    for sev in severity_order:
        count = st["counts"].get(sev, 0)
        lines.append(f"| {sev.capitalize()} | {count} |")
    lines.append("")

    if report.methodology:
        lines.append("## Methodology")
        lines.append("")
        lines.append(report.methodology)
        lines.append("")

    lines.append("## Findings Detail")
    lines.append("")
    if findings:
        for sev in severity_order:
            items = [f for f in findings if f.severity == sev]
            if not items:
                continue
            for f in sorted(items, key=lambda x: x.title.lower()):
                lines.append(f"### {f.title} `{sev}`")
                lines.append("")
                if f.cvss or f.cwe or f.cve:
                    lines.append("| Reference | Value |")
                    lines.append("|---|---|")
                    if f.cvss:
                        lines.append(f"| CVSS | {f.cvss} |")
                    if f.cwe:
                        lines.append(f"| CWE | CWE-{f.cwe} |")
                    if f.cve:
                        lines.append(f"| CVE | {f.cve} |")
                    lines.append("")
                if f.asset:
                    lines.append(f"**Asset:** {f.asset}")
                    lines.append("")
                if f.description:
                    lines.append(f"**Description.** {f.description}")
                    lines.append("")
                if f.impact:
                    lines.append(f"**Impact.** {f.impact}")
                    lines.append("")
                if f.evidence:
                    lines.append("**Evidence:**")
                    lines.append("")
                    lines.append("```")
                    lines.append(f.evidence)
                    lines.append("```")
                    lines.append("")
                if f.recommendation:
                    lines.append(f"**Recommendation.** {f.recommendation}")
                    lines.append("")
    else:
        lines.append("*No findings recorded.*")
        lines.append("")

    return "\n".join(lines)